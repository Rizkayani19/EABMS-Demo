import io
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from core.models import Transaction, Unit, InternalSection


def _money(v):
    return Decimal(str(v or 0)).quantize(Decimal('0.01'))


def _name(v):
    return getattr(v, 'name', '') if v else ''


def _excel_value(v):
    return v.replace(tzinfo=None) if isinstance(v, datetime) and v.tzinfo is not None else v


def _clean_table_name(name):
    clean = re.sub(r'[^A-Za-z0-9_]', '', str(name or 'Table'))
    if not clean or clean[0].isdigit():
        clean = 'T_' + clean
    return clean[:250]


def _style(ws, header_row=10, table_name=None):
    for cell in ws[header_row]:
        cell.value = str(cell.value or '')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1E293B', end_color='1E293B', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws.freeze_panes = f'A{header_row + 1}'

    if table_name and ws.max_row > header_row:
        safe_name = _clean_table_name(table_name)
        ref = f'A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}'
        tab = Table(displayName=safe_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(
            name='TableStyleMedium2',
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tab)

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)


def _filter_label(filters, key, model):
    ids = filters.get(key) or []
    if not ids:
        return 'All'
    return ', '.join(
        model.objects.filter(pk__in=ids).order_by('display_order', 'name').values_list('name', flat=True)
    ) or 'All'


def _metadata(ws, filters, title):
    ws.append(['ABMS — Enterprise Accounts & Business Management System (EABMS)'])
    ws.append([str(title)])
    ws.append([f"Account: {_name(filters.get('account')) or 'All'}"])
    ws.append([f"Financial Year: {_name(filters.get('financial_year')) or 'All'}"])
    start, end = filters.get('start_date'), filters.get('end_date')
    ws.append([f"From Date: {start:%d-%m-%Y}" if start else 'From Date: Start of selected scope'])
    ws.append([f"To Date: {end:%d-%m-%Y}" if end else 'To Date: End of selected scope'])
    ws.append([f"Unit(s): {_filter_label(filters, 'unit_ids', Unit)}"])
    ws.append([f"Internal Section(s): {_filter_label(filters, 'internal_section_ids', InternalSection)}"])
    ws.append([])


def _filtered_transactions(filters):
    qs = Transaction.objects.filter(status=Transaction.Status.ACTIVE)
    if filters.get('account'):
        qs = qs.filter(account=filters['account'])
    if filters.get('financial_year'):
        qs = qs.filter(financial_year=filters['financial_year'])
    if filters.get('start_date'):
        qs = qs.filter(transaction_date__gte=filters['start_date'])
    if filters.get('end_date'):
        qs = qs.filter(transaction_date__lte=filters['end_date'])
    if filters.get('direction'):
        qs = qs.filter(direction=filters['direction'])
    if filters.get('major_head'):
        qs = qs.filter(major_head=filters['major_head'])
    if filters.get('account_head'):
        qs = qs.filter(account_head=filters['account_head'])
    if filters.get('sub_head'):
        qs = qs.filter(sub_head=filters['sub_head'])
    if filters.get('outflow_nature'):
        qs = qs.filter(outflow_nature=filters['outflow_nature'])
    if filters.get('unit_ids'):
        qs = qs.filter(transaction_units__id__in=filters['unit_ids'])
    if filters.get('internal_section_ids'):
        qs = qs.filter(transaction_sections__id__in=filters['internal_section_ids'])
    return qs.distinct().select_related(
        'account', 'financial_year', 'major_head', 'account_head', 'sub_head',
        'outflow_nature', 'unit', 'internal_section'
    ).prefetch_related('transaction_units', 'transaction_sections').order_by('transaction_date', 'pk')


def _append_summary(wb, filters, title, transactions, key_func, table_name):
    safe_sheet_title = title[:31]
    ws = wb.create_sheet(safe_sheet_title)
    _metadata(ws, filters, title)
    ws.append(['Category', 'Inflows', 'Outflows', 'Net Movement'])
    data = defaultdict(lambda: [Decimal('0.00'), Decimal('0.00')])
    for tx in transactions:
        key = key_func(tx) or 'Unclassified'
        data[key][0 if tx.direction == Transaction.Direction.INFLOW else 1] += _money(tx.amount)
    for key in sorted(data, key=str):
        i, o = data[key]
        ws.append([str(key), i, o, i - o])
    _style(ws, header_row=10, table_name=table_name)


def export_management_workbook(filters):
    transactions = list(_filtered_transactions(filters))
    wb = Workbook()

    # 1. Executive Summary
    ws = wb.active
    ws.title = 'Executive Summary'
    _metadata(ws, filters, 'Executive Summary')
    total_in = sum((_money(t.amount) for t in transactions if t.direction == Transaction.Direction.INFLOW), Decimal('0.00'))
    total_out = sum((_money(t.amount) for t in transactions if t.direction == Transaction.Direction.OUTFLOW), Decimal('0.00'))
    ws.append(['Metric', 'Amount'])
    ws.append(['Total Transactions Recorded', len(transactions)])
    ws.append(['Total Inflows (Actual)', total_in])
    ws.append(['Total Outflows (Actual)', total_out])
    ws.append(['Net Cash Movement', total_in - total_out])
    _style(ws, header_row=10, table_name='ExecSummary')

    # 2. Report Parameters
    ws = wb.create_sheet('Report Parameters')
    _metadata(ws, filters, 'Report Parameters')
    ws.append(['Parameter', 'Applied Filter Value'])
    ws.append(['Account', _name(filters.get('account')) or 'All Accounts'])
    ws.append(['Financial Year', _name(filters.get('financial_year')) or 'All Financial Years'])
    ws.append(['From Date', f"{filters.get('start_date'):%d-%m-%Y}" if filters.get('start_date') else 'Start of Scope'])
    ws.append(['To Date', f"{filters.get('end_date'):%d-%m-%Y}" if filters.get('end_date') else 'End of Scope'])
    ws.append(['Unit(s)', _filter_label(filters, 'unit_ids', Unit)])
    ws.append(['Internal Section(s)', _filter_label(filters, 'internal_section_ids', InternalSection)])
    _style(ws, header_row=10, table_name='ReportParams')

    # 3. Master Transactions
    ws = wb.create_sheet('Master Transactions')
    _metadata(ws, filters, 'Master Transactions — Single Source of Truth Export')
    headers = [
        'Transaction ID', 'Account', 'Financial Year', 'Date', 'Month',
        'Direction', 'Description', 'Amount', 'Major Head', 'Account Head',
        'Sub Head', 'Outflow Nature', 'Unit(s)', 'Internal Section(s)',
        'Reference', 'Document Reference', 'Created By', 'Created At', 'Status'
    ]
    ws.append(headers)
    for tx in transactions:
        units = sorted({u.name for u in tx.transaction_units.all()}) or ([tx.unit.name] if tx.unit else [])
        sections = sorted({s.name for s in tx.transaction_sections.all()}) or ([tx.internal_section.name] if tx.internal_section else [])
        ws.append([
            str(tx.transaction_number),
            _name(tx.account),
            _name(tx.financial_year),
            tx.transaction_date,
            tx.transaction_date.strftime('%Y-%m'),
            str(tx.direction),
            str(tx.description or tx.remarks or ''),
            _money(tx.amount),
            _name(tx.major_head),
            _name(tx.account_head),
            _name(tx.sub_head),
            _name(tx.outflow_nature),
            ', '.join(units),
            ', '.join(sections),
            str(tx.reference or ''),
            str(tx.document_reference or ''),
            _name(tx.created_by),
            _excel_value(tx.created_at),
            str(tx.status)
        ])
    _style(ws, header_row=10, table_name='MasterTx')

    # 4. Inflows Summary
    _append_summary(
        wb, filters, 'Inflows Summary',
        [t for t in transactions if t.direction == Transaction.Direction.INFLOW],
        lambda t: _name(t.account_head) or t.transaction_number,
        'InflowsSummary'
    )

    # 5. Outflows Summary
    _append_summary(
        wb, filters, 'Outflows Summary',
        [t for t in transactions if t.direction == Transaction.Direction.OUTFLOW],
        lambda t: _name(t.account_head) or t.transaction_number,
        'OutflowsSummary'
    )

    # 6. Monthly Summary
    _append_summary(
        wb, filters, 'Monthly Summary',
        transactions,
        lambda t: t.transaction_date.strftime('%Y-%m'),
        'MonthlySummary'
    )

    # 7. Account Head Summary
    _append_summary(
        wb, filters, 'Account Head Summary',
        transactions,
        lambda t: _name(t.account_head),
        'AccountHeadSummary'
    )

    # 8. Sub Head Summary
    _append_summary(
        wb, filters, 'Sub Head Summary',
        transactions,
        lambda t: _name(t.sub_head),
        'SubHeadSummary'
    )

    # 9. Unit Summary
    def unit_selector(t):
        return {u.name for u in t.transaction_units.all()} or ({t.unit.name} if t.unit else set())

    ws_unit = wb.create_sheet('Unit Summary')
    _metadata(ws_unit, filters, 'Unit Summary')
    ws_unit.append(['Unit', 'Inflows', 'Outflows', 'Net Movement'])
    unit_data = defaultdict(lambda: [Decimal('0.00'), Decimal('0.00')])
    for tx in transactions:
        for val in unit_selector(tx):
            unit_data[val][0 if tx.direction == Transaction.Direction.INFLOW else 1] += _money(tx.amount)
    for key in sorted(unit_data, key=str):
        i, o = unit_data[key]
        ws_unit.append([str(key), i, o, i - o])
    _style(ws_unit, header_row=10, table_name='UnitSummary')

    # 10. Internal Section Summary
    def section_selector(t):
        return {s.name for s in t.transaction_sections.all()} or ({t.internal_section.name} if t.internal_section else set())

    ws_sec = wb.create_sheet('Internal Section Summary'[:31])
    _metadata(ws_sec, filters, 'Internal Section Summary')
    ws_sec.append(['Internal Section', 'Inflows', 'Outflows', 'Net Movement'])
    sec_data = defaultdict(lambda: [Decimal('0.00'), Decimal('0.00')])
    for tx in transactions:
        for val in section_selector(tx):
            sec_data[val][0 if tx.direction == Transaction.Direction.INFLOW else 1] += _money(tx.amount)
    for key in sorted(sec_data, key=str):
        i, o = sec_data[key]
        ws_sec.append([str(key), i, o, i - o])
    _style(ws_sec, header_row=10, table_name='InternalSectionSummary')

    # 11. Unit + Section Analysis
    ws_us = wb.create_sheet('Unit + Section Analysis'[:31])
    _metadata(ws_us, filters, 'Unit + Internal Section Analysis')
    ws_us.append(['Unit', 'Internal Section', 'Inflows', 'Outflows', 'Net Movement'])
    us_data = defaultdict(lambda: [Decimal('0.00'), Decimal('0.00')])
    for tx in transactions:
        units = list(tx.transaction_units.all()) or ([tx.unit] if tx.unit else [None])
        sections = list(tx.transaction_sections.all()) or ([tx.internal_section] if tx.internal_section else [None])
        for u in units:
            for s in sections:
                key = (_name(u) or 'Unspecified Unit', _name(s) or 'Unspecified Section')
                us_data[key][0 if tx.direction == Transaction.Direction.INFLOW else 1] += _money(tx.amount)
    for (u, s), (i, o) in sorted(us_data.items()):
        ws_us.append([str(u), str(s), i, o, i - o])
    _style(ws_us, header_row=10, table_name='UnitSectionAnalysis')

    # 12. Deleted Transactions (Audit Log)
    ws_del = wb.create_sheet('Deleted Transactions'[:31])
    _metadata(ws_del, filters, 'Deleted Transactions — Authorised Audit View')
    ws_del.append(['Transaction ID', 'Date', 'Amount', 'Deleted By', 'Deleted At', 'Deletion Reason'])
    deleted = Transaction.objects.filter(status=Transaction.Status.SOFT_DELETED)
    if filters.get('account'):
        deleted = deleted.filter(account=filters['account'])
    if filters.get('financial_year'):
        deleted = deleted.filter(financial_year=filters['financial_year'])
    if filters.get('start_date'):
        deleted = deleted.filter(transaction_date__gte=filters['start_date'])
    if filters.get('end_date'):
        deleted = deleted.filter(transaction_date__lte=filters['end_date'])
    for tx in deleted.select_related('deleted_by'):
        ws_del.append([
            str(tx.transaction_number),
            tx.transaction_date,
            _money(tx.amount),
            _name(tx.deleted_by),
            _excel_value(tx.deleted_at),
            str(tx.deletion_reason or '')
        ])
    _style(ws_del, header_row=10, table_name='DeletedTx')

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
