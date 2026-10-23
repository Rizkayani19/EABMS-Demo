"""Excel exports for existing-budget analysis and future-budget planning."""
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from core.budget_analysis_services import analyse_budget, variance_flags
from core.future_budget_services import get_future_budget_summary
from core.future_budget_allocation_services import allocation_reconciliation, build_future_organisational_analysis
from core.future_budget_workbook_contract import build_calculation_detail_row, build_workbook_controls
from core.models import FutureBudget


def _money(v):
    return Decimal(str(v or 0)).quantize(Decimal("0.01"))


def _style(ws, header_row, table_name=None):
    for c in ws[header_row]:
        c.font = Font(bold=True)
    ws.freeze_panes = f"A{header_row + 1}"
    if table_name and ws.max_row > header_row:
        ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
        tab = Table(displayName=table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    for col in ws.columns:
        width = min(max(len(str(c.value or "")) for c in col) + 2, 50)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(width, 10)


def _header(ws, title, account, source_fy, target_fy=None, units=None, internal_sections=None):
    ws.append(["ABMS — Enterprise Accounts & Business Management System (EABMS)"])
    ws.append([title])
    ws.append([f"Account: {getattr(account, 'name', 'All')}"])
    ws.append([f"Source Financial Year: {getattr(source_fy, 'name', 'All')}"])
    if target_fy:
        ws.append([f"Target Financial Year: {getattr(target_fy, 'name', 'All')}"])
    ws.append([f"Unit Filter: {', '.join(units) if units else 'All'}"])
    ws.append([f"Internal Section Filter: {', '.join(internal_sections) if internal_sections else 'All'}"])
    ws.append([])


def _filter_org_analysis(org_analysis, units=None, internal_sections=None):
    units = set(units or [])
    sections = set(internal_sections or [])
    if not units and not sections:
        return org_analysis
    filtered = dict(org_analysis)
    filtered['unit_rows'] = [r for r in org_analysis['unit_rows'] if not units or r['unit'] in units]
    filtered['section_rows'] = [r for r in org_analysis['section_rows'] if not sections or r['internal_section'] in sections]
    filtered['matrix_rows'] = [r for r in org_analysis['matrix_rows'] if (not units or r['unit'] in units) and (not sections or r['internal_section'] in sections)]
    filtered['head_unit_rows'] = [r for r in org_analysis['head_unit_rows'] if not units or r['unit'] in units]
    filtered['head_section_rows'] = [r for r in org_analysis['head_section_rows'] if not sections or r['internal_section'] in sections]
    filtered['future_allocation_total'] = sum((r['future_allocated'] for r in filtered['unit_rows']), Decimal('0.00'))
    filtered['previous_unit_actual_total'] = sum((r['previous_actual'] for r in filtered['unit_rows']), Decimal('0.00'))
    filtered['previous_section_actual_total'] = sum((r['previous_actual'] for r in filtered['section_rows']), Decimal('0.00'))
    return filtered


def export_budget_planning_workbook(account, source_fy, target_fy=None, units=None, internal_sections=None):
    """Build an auditable management workbook from canonical backend results."""
    analysis = analyse_budget(account, source_fy)
    flags = variance_flags(analysis)
    future = None
    future_summary = None
    org_analysis = None
    if target_fy:
        future = FutureBudget.objects.filter(account=account, source_fy=source_fy, target_fy=target_fy).first()
        if future:
            future_summary = get_future_budget_summary(future)
            org_analysis = _filter_org_analysis(build_future_organisational_analysis(future), units, internal_sections)

    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    _header(ws, "Budget Planning Executive Summary", account, source_fy, target_fy, units, internal_sections)
    ws.append(["Metric", "Amount"])
    ws.append(["Previous Budget", _money(analysis["total_budget"])])
    ws.append(["Previous Actual", _money(analysis["total_actual"])])
    ws.append(["Budget Variance", _money(analysis["total_variance"])])
    ws.append(["Actual Utilisation %", _money(analysis["utilisation_pct"])])
    if future_summary:
        ws.append(["Future System Suggested", _money(future_summary["system_suggested_total"])])
        ws.append(["Future Management Proposed", _money(future_summary["user_proposed_total"])])
        ws.append(["Future Approved", _money(future_summary["approved_total"])])
        ws.append(["Proposed vs Previous Actual", _money(future_summary["diff_from_prev_actual"])])
        ws.append(["Proposed vs Previous Budget", _money(future_summary["diff_from_prev_budget"])])
        if org_analysis:
            ws.append(["Future Organisational Allocation (filtered)", _money(org_analysis["future_allocation_total"])])
            ws.append(["Unreconciled Future Budget Lines", org_analysis["unreconciled_lines"]])
    _style(ws, 8)

    ws = wb.create_sheet("Existing Budget Analysis")
    _header(ws, "Existing Budget — Budget vs Actual Analysis", account, source_fy)
    ws.append(["Account Head", "Previous Budget", "Actual", "Variance", "Utilisation %", "Status"])
    for r in analysis["rows"]:
        ws.append([r["account_head"], _money(r["budget"]), _money(r["actual"]), _money(r["variance"]), _money(r["utilisation_pct"]), r["status"]])
    _style(ws, 8, "ExistingBudgetAnalysis")

    ws = wb.create_sheet("Material Variances")
    _header(ws, "Material Budget Variance Flags", account, source_fy)
    ws.append(["Account Head", "Budget", "Actual", "Utilisation %", "Flag"])
    for r in flags:
        ws.append([r["account_head"], _money(r["budget"]), _money(r["actual"]), _money(r["utilisation_pct"]), r["flag"]])
    _style(ws, 8, "MaterialVariances")

    if future:
        ws = wb.create_sheet("Future Budget Lines")
        _header(ws, "Future Budget Proposal — Line Analysis", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Account Head", "Previous Budget", "Previous Actual", "Routine", "Known Changes", "One-time", "System Suggested", "Management Proposed", "Approved", "Continuing Flag"])
        for line in future.lines.select_related("account_head").order_by("account_head__name"):
            ws.append([line.account_head.name, _money(line.previous_budget), _money(line.previous_actual), _money(line.routine_amount), _money(line.known_change_amount), _money(line.onetime_amount), _money(line.system_suggested), _money(line.user_proposed), _money(line.approved_amount), line.continuing_flag or ""])
        _style(ws, 8, "FutureBudgetLines")

        ws = wb.create_sheet("Calculation Detail")
        _header(ws, "Future Budget — Canonical Calculation Detail", account, source_fy, target_fy, units, internal_sections)
        headers = ["Account Head", "Recurring Baseline", "Routine Adjustment", "Continuing Requirement", "Management Additional", "Known Changes", "Personnel Effect", "One-Time / Discontinued", "System Suggested", "Management Proposed", "Approved"]
        ws.append(headers)
        for line in future.lines.select_related("account_head").order_by("account_head__name"):
            row = build_calculation_detail_row(line)
            ws.append([row[h] for h in headers])
        _style(ws, 8, "FutureCalculationDetail")

        ws = wb.create_sheet("Management Inputs")
        _header(ws, "Future Budget — Management Inputs", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Account Head", "Input Type", "Description", "Amount", "Status", "Reference"])
        for line in future.lines.select_related("account_head"):
            for detail in (line.suggestion_breakdown or {}).get("management_input_detail", []):
                ws.append([line.account_head.name, detail.get("type", ""), detail.get("description", ""), _money(detail.get("amount")), "ACTIVE", detail.get("reference", "")])
        _style(ws, 8, "FutureManagementInputs")

        ws = wb.create_sheet("Controls")
        _header(ws, "Future Budget — Controls", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Control", "Value", "Status", "Detail"])
        for control in build_workbook_controls(future):
            ws.append([control["control"], control["value"], control["status"], control.get("detail", "")])
        _style(ws, 8, "FutureBudgetControls")

        ws = wb.create_sheet("Future Organisational Allocation")
        _header(ws, "Future Budget — Unit / Internal Section Monetary Allocation", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Account Head", "Unit", "Internal Section", "Allocated Amount", "Justification", "Remarks"])
        allocations = future.lines.select_related("account_head").prefetch_related("allocations__unit", "allocations__internal_section")
        for line in allocations:
            for a in line.allocations.all().order_by("unit__name", "internal_section__name"):
                unit_name = a.unit.name if a.unit else "All Units"
                section_name = a.internal_section.name if a.internal_section else "All Sections"
                if units and a.unit is not None and unit_name not in units:
                    continue
                if internal_sections and a.internal_section is not None and section_name not in internal_sections:
                    continue
                ws.append([line.account_head.name, unit_name, section_name, _money(a.allocated_amount), a.justification, a.remarks])
        _style(ws, 8, "FutureOrgAllocation")

        ws = wb.create_sheet("Allocation Reconciliation")
        _header(ws, "Future Budget — Allocation Reconciliation", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Account Head", "Target Amount", "Allocated Amount", "Difference", "Status", "Allocation Rows"])
        for line in future.lines.select_related("account_head").prefetch_related("allocations"):
            rec = allocation_reconciliation(line)
            ws.append([line.account_head.name, _money(rec["target"]), _money(rec["allocated"]), _money(rec["difference"]), "RECONCILED" if rec["is_reconciled"] else "NOT RECONCILED", line.allocations.count()])
        _style(ws, 8, "AllocationReconciliation")

        ws = wb.create_sheet("Allocation Summary")
        _header(ws, "Future Budget — Organisation Summary", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Dimension", "Organisation", "Previous Actual", "Future Allocated", "Difference", "Change %"])
        for r in org_analysis["unit_rows"]:
            ws.append(["Unit", r["unit"], _money(r["previous_actual"]), _money(r["future_allocated"]), _money(r["difference"]), r["change_pct"]])
        for r in org_analysis["section_rows"]:
            ws.append(["Internal Section", r["internal_section"], _money(r["previous_actual"]), _money(r["future_allocated"]), _money(r["difference"]), r["change_pct"]])
        _style(ws, 8, "FutureOrgSummary")

        ws = wb.create_sheet("Unit-Section Matrix")
        _header(ws, "Future Budget — Unit + Internal Section Analysis", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Unit", "Internal Section", "Future Allocated"])
        for r in org_analysis["matrix_rows"]:
            ws.append([r["unit"], r["internal_section"], _money(r["future_allocated"])])
        _style(ws, 8, "UnitSectionMatrix")

        ws = wb.create_sheet("Head-Unit Analysis")
        _header(ws, "Future Budget — Account Head + Unit Analysis", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Account Head", "Unit", "Future Allocated"])
        for r in org_analysis["head_unit_rows"]:
            ws.append([r["account_head"], r["unit"], _money(r["future_allocated"])])
        _style(ws, 8, "HeadUnitAnalysis")

        ws = wb.create_sheet("Head-Section Analysis")
        _header(ws, "Future Budget — Account Head + Internal Section Analysis", account, source_fy, target_fy, units, internal_sections)
        ws.append(["Account Head", "Internal Section", "Future Allocated"])
        for r in org_analysis["head_section_rows"]:
            ws.append([r["account_head"], r["internal_section"], _money(r["future_allocated"])])
        _style(ws, 8, "HeadSectionAnalysis")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
