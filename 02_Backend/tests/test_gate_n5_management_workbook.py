from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from core.models import Account, FinancialYear, AccountHead, Transaction, Unit, InternalSection
from reports.exporters.management_workbook import export_management_workbook


class GateN5ManagementWorkbookTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        account = Account.objects.create(code="N5", name="N5 Test Account")
        fy = FinancialYear.objects.create(code="N5-26-27", name="FY 2026-27 N5", start_date=date(2026, 7, 1), end_date=date(2027, 6, 30))
        head = AccountHead.objects.create(account=account, name="N5 Head")
        unit_a = Unit.objects.create(code="N5-A", name="N5 Unit A")
        unit_b = Unit.objects.create(code="N5-B", name="N5 Unit B")
        section = InternalSection.objects.create(code="N5-EVT", name="N5 EVT")
        tx = Transaction.objects.create(transaction_number="N5-TX-001", account=account, financial_year=fy, account_head=head, transaction_date=date(2026, 7, 15), direction=Transaction.Direction.OUTFLOW, amount=Decimal("1000.00"), description="N5 allocation test", status=Transaction.Status.ACTIVE)
        tx.transaction_units.add(unit_a, unit_b)
        tx.transaction_sections.add(section)
        cls.account, cls.fy = account, fy
        cls.unit_a, cls.section = unit_a, section

    def test_workbook_contains_transaction_and_current_allocation_grain(self):
        data = export_management_workbook({"account": self.account, "financial_year": self.fy})
        wb = load_workbook(BytesIO(data), data_only=False)
        self.assertIn("Master Transactions", wb.sheetnames)
        self.assertIn("Unit Summary", wb.sheetnames)
        self.assertIn("Internal Section Summary", wb.sheetnames)
        self.assertIn("Unit + Section Analysis", wb.sheetnames)
        self.assertEqual(wb["Master Transactions"].max_row, 11)

    def test_unit_filter_uses_allocation_scope_without_double_counting(self):
        data = export_management_workbook({"account": self.account, "financial_year": self.fy, "unit_ids": [self.unit_a.pk], "internal_section_ids": [self.section.pk]})
        wb = load_workbook(BytesIO(data), data_only=False)
        ws = wb["Unit + Section Analysis"]
        rows = [r for r in ws.iter_rows(min_row=10, values_only=True) if r[0] == "N5 Unit A"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][3], Decimal("1000.00"))
