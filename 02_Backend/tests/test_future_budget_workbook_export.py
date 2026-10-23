from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook
from django.contrib.auth import get_user_model

from core.budget_reporting_export import export_budget_planning_workbook
from core.models import Account, AccountHead, FinancialYear, FutureBudget, FutureBudgetLine


@pytest.mark.django_db
def test_future_budget_workbook_contains_canonical_sheets_and_fields():
    account = Account.objects.create(code="G23WB", name="G23 Workbook Account")
    source = FinancialYear.objects.create(
        code="2041-42-wb",
        name="2041-42",
        start_date=date(2041, 7, 1),
        end_date=date(2042, 6, 30),
    )
    target = FinancialYear.objects.create(
        code="2043-44-wb",
        name="2043-44",
        start_date=date(2043, 7, 1),
        end_date=date(2044, 6, 30),
    )
    head = AccountHead.objects.create(name="G23 Workbook Head", account=account)
    budget = FutureBudget.objects.create(account=account, source_fy=source, target_fy=target)
    FutureBudgetLine.objects.create(
        future_budget=budget,
        account_head=head,
        previous_budget=Decimal("900.00"),
        previous_actual=Decimal("1000.00"),
        system_suggested=Decimal("1100.00"),
        user_proposed=Decimal("1050.00"),
        approved_amount=Decimal("1050.00"),
        suggestion_breakdown={
            "previous_actual": "1000.00",
            "routine_adjustment": "100.00",
            "routine_adjustment_pct": "10.00",
            "continuing_requirement": "0.00",
            "management_additional_requirement": "0.00",
            "known_changes": "0.00",
            "personnel_effect": "0.00",
            "onetime_deduction": "0.00",
            "system_suggested": "1100.00",
        },
    )

    data = export_budget_planning_workbook(account, source, target)
    wb = load_workbook(BytesIO(data), data_only=False)

    assert "Future Budget Lines" in wb.sheetnames
    assert "Calculation Detail" in wb.sheetnames
    assert "Management Inputs" in wb.sheetnames
    assert "Controls" in wb.sheetnames
    ws = wb["Calculation Detail"]
    headers = [c.value for c in ws[9]]
    assert "Recurring Baseline" in headers
    assert "Routine Adjustment" in headers
    assert "Management Additional" in headers
    assert "Personnel Effect" in headers
    assert "System Suggested" in headers
    assert "Management Proposed" in headers
    assert "Approved" in headers


@pytest.mark.django_db
def test_workbook_is_reopenable_and_uses_backend_values():
    account = Account.objects.create(code="G23RO", name="G23 Workbook Reopen")
    source = FinancialYear.objects.create(
        code="2043-44-ro",
        name="2043-44",
        start_date=date(2043, 7, 1),
        end_date=date(2044, 6, 30),
    )
    target = FinancialYear.objects.create(
        code="2044-45-ro",
        name="2044-45",
        start_date=date(2044, 7, 1),
        end_date=date(2045, 6, 30),
    )
    head = AccountHead.objects.create(name="G23 Reopen Head", account=account)
    budget = FutureBudget.objects.create(account=account, source_fy=source, target_fy=target)
    FutureBudgetLine.objects.create(future_budget=budget, account_head=head, system_suggested=Decimal("1234.56"))

    data = export_budget_planning_workbook(account, source, target)
    wb = load_workbook(BytesIO(data), data_only=False)
    ws = wb["Future Budget Lines"]
    values = [c.value for c in ws[10]]
    assert Decimal(str(values[6])) == Decimal("1234.56")
