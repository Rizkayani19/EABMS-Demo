from datetime import date
from decimal import Decimal
from io import BytesIO
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import JsonResponse, HttpResponse
from django.shortcuts import render
from django.views import View
from django.db.models import Sum

from core.models import Account, FinancialYear, AccountHead, SubHead, Unit, InternalSection, Transaction
from reports.services import build_budget_vs_actual, build_monthly_outflows, build_account_head_summary
from reports.organisational_v2 import build_unit_summary, build_internal_section_summary, build_unit_section_summary
from reports.exporters.management_workbook import export_management_workbook


def _ids(request, key):
    values = request.GET.getlist(key) or request.GET.getlist(f'{key}[]')
    return [int(v) for v in values if str(v).isdigit()]


def build_report_filters(request):
    account = None
    fy = None
    if request.GET.get('account', '').isdigit(): account = Account.objects.filter(pk=request.GET['account']).first()
    if request.GET.get('financial_year', '').isdigit(): fy = FinancialYear.objects.filter(pk=request.GET['financial_year']).first()
    account = account or getattr(request, 'account', None)
    fy = fy or getattr(request, 'fy', None)
    start = request.GET.get('date_from') or request.GET.get('start_date')
    end = request.GET.get('date_to') or request.GET.get('end_date')
    try: start = date.fromisoformat(start) if start else None
    except ValueError: start = None
    try: end = date.fromisoformat(end) if end else None
    except ValueError: end = None
    return {'account': account, 'financial_year': fy, 'start_date': start, 'end_date': end,
            'unit_ids': list(dict.fromkeys(_ids(request, 'unit') + _ids(request, 'unit_ids'))),
            'internal_section_ids': list(dict.fromkeys(_ids(request, 'internal_section') + _ids(request, 'internal_section_ids'))),
            'direction': request.GET.get('direction') or None, 'major_head': None, 'sub_head': None,
            'account_head': None, 'outflow_nature': None}


def _org_filter_transactions(filters):
    qs = Transaction.objects.filter(status=Transaction.Status.ACTIVE)
    if filters.get('account'): qs = qs.filter(account=filters['account'])
    if filters.get('financial_year'): qs = qs.filter(financial_year=filters['financial_year'])
    if filters.get('start_date'): qs = qs.filter(transaction_date__gte=filters['start_date'])
    if filters.get('end_date'): qs = qs.filter(transaction_date__lte=filters['end_date'])
    if filters.get('direction'): qs = qs.filter(direction=filters['direction'])
    if filters.get('unit_ids'): qs = qs.filter(transaction_units__id__in=filters['unit_ids'])
    if filters.get('internal_section_ids'): qs = qs.filter(transaction_sections__id__in=filters['internal_section_ids'])
    return qs.distinct()


class ReportCenterView(LoginRequiredMixin, View):
    template_name = 'reports/report_center.html'
    def get(self, request):
        return render(request, self.template_name, {'accounts': Account.objects.all().order_by('name'), 'financial_years': FinancialYear.objects.all().order_by('-start_date'), 'units': Unit.objects.filter(is_active=True).order_by('name'), 'internal_sections': InternalSection.objects.filter(is_active=True).order_by('display_order', 'name'), 'account_heads': AccountHead.objects.all().order_by('name'), 'sub_heads': SubHead.objects.all().order_by('name'), 'account': getattr(request, 'account', None), 'fy': getattr(request, 'fy', None)})


class ReportCenterDataView(LoginRequiredMixin, View):
    def get(self, request):
        filters = build_report_filters(request)
        tx_qs = _org_filter_transactions(filters)
        totals = tx_qs.values('direction').annotate(total=Sum('amount'))
        inflows = next((x['total'] for x in totals if x['direction'] == 'INFLOW'), Decimal('0.00'))
        outflows = next((x['total'] for x in totals if x['direction'] == 'OUTFLOW'), Decimal('0.00'))
        result = {'filters': {'account': getattr(filters.get('account'), 'name', 'All'), 'financial_year': getattr(filters.get('financial_year'), 'name', 'All'), 'date_from': filters['start_date'].isoformat() if filters['start_date'] else '', 'date_to': filters['end_date'].isoformat() if filters['end_date'] else '', 'units': list(Unit.objects.filter(pk__in=filters['unit_ids']).values_list('name', flat=True)), 'sections': list(InternalSection.objects.filter(pk__in=filters['internal_section_ids']).values_list('name', flat=True))}, 'summary': {'transaction_count': tx_qs.count(), 'inflows': str(inflows), 'outflows': str(outflows), 'net': str(inflows - outflows)}, 'unit_summary': [{'unit': r['unit'], 'amount': str(r['amount'])} for r in build_unit_summary(filters)], 'section_summary': [{'internal_section': r['internal_section'], 'amount': str(r['amount'])} for r in build_internal_section_summary(filters)], 'unit_section_summary': [{'unit': r['unit'], 'internal_section': r['internal_section'], 'amount': str(r['amount'])} for r in build_unit_section_summary(filters)]}
        if filters.get('account') and filters.get('financial_year'):
            result['budget_vs_actual'] = [{k: str(v) if isinstance(v, Decimal) else v for k, v in row.items()} for row in build_budget_vs_actual(filters)]
            result['monthly_outflows'] = [{k: str(v) if isinstance(v, Decimal) else v for k, v in row.items()} for row in build_monthly_outflows(filters)]
            result['account_head_summary'] = [{k: str(v) if isinstance(v, Decimal) else v for k, v in row.items()} for row in build_account_head_summary(filters)]
        return JsonResponse(result)


class ManagementWorkbookExportView(LoginRequiredMixin, View):
    def get(self, request):
        filters = build_report_filters(request)
        payload = export_management_workbook(filters)
        profile = getattr(request.user, 'profile', None)
        if not profile or profile.role != 'ADMINISTRATOR':
            # Deleted Transactions is an Administrator-only audit view.
            from openpyxl import load_workbook
            wb = load_workbook(BytesIO(payload))
            if 'Deleted Transactions' in wb.sheetnames:
                del wb['Deleted Transactions']
            out = BytesIO(); wb.save(out); payload = out.getvalue()
        response = HttpResponse(payload, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="ABMS_Management_Report.xlsx"'
        return response
