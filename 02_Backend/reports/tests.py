from datetime import date
from decimal import Decimal
from types import SimpleNamespace

from django.test import SimpleTestCase
from openpyxl import load_workbook

from .exporters.excel import _pakistan_now_label, export_to_excel
from .exporters.pdf import _pakistan_now_label as pdf_pakistan_now_label, export_to_pdf
from .views import _trim_month_rows_to_selected_range


class ReportDateAndExportTests(SimpleTestCase):
    def setUp(self):
        self.report = SimpleNamespace(
            title="Test Report",
            slug="test-report",
            description="Test",
            columns=[{"key": "month", "label": "Month"}, {"key": "amount", "label": "Amount"}],
            totals=["amount"],
        )

    def test_custom_range_removes_months_outside_selected_range(self):
        dataset = [
            {"month": "2026-07", "amount": Decimal("10.00")},
            {"month": "2026-08", "amount": Decimal("20.00")},
            {"month": "2026-09", "amount": Decimal("30.00")},
        ]
        filters = {"start_date": date(2026, 8, 15), "end_date": date(2026, 8, 20)}
        result = _trim_month_rows_to_selected_range(dataset, filters)
        self.assertEqual([row["month"] for row in result], ["2026-08"])

    def test_month_rows_are_unchanged_without_custom_range(self):
        dataset = [{"month": "2026-07", "amount": Decimal("10.00")}, {"month": "2026-08", "amount": Decimal("20.00")}]
        self.assertEqual(_trim_month_rows_to_selected_range(dataset, {}), dataset)

    def test_excel_export_records_pakistan_time(self):
        payload = export_to_excel(self.report, [{"month": "2026-08", "amount": Decimal("20.00")}], {})
        workbook = load_workbook(filename=__import__("io").BytesIO(payload), read_only=True)
        values = [cell.value for row in workbook.active.iter_rows() for cell in row if cell.value]
        self.assertTrue(any("Generated:" in str(value) and " PKT (UTC+05:00)" in str(value) for value in values))

    def test_timestamp_helpers_use_pakistan_standard_time(self):
        excel_label = _pakistan_now_label()
        pdf_label = pdf_pakistan_now_label()
        self.assertRegex(excel_label, r"^\d{2}-\d{2}-\d{4} \d{2}:\d{2} PKT \(UTC\+05:00\)$")
        self.assertRegex(pdf_label, r"^\d{2}-\d{2}-\d{4} \d{2}:\d{2} PKT \(UTC\+05:00\)$")

    def test_pdf_export_completes_with_selected_period(self):
        payload = export_to_pdf(
            self.report,
            [{"month": "2026-08", "amount": Decimal("20.00")}],
            {"start_date": date(2026, 8, 1), "end_date": date(2026, 8, 31)},
        )
        self.assertTrue(payload.startswith(b"%PDF-"))
