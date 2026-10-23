"""Management-grade Future Budget workbook built on the existing SSOT exporter."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from core.budget_reporting_export import export_budget_planning_workbook
from core.future_budget_services import get_future_budget_summary
from core.future_budget_allocation_services import build_future_organisational_analysis
from core.future_budget_management_models import FutureBudgetManagementInput
from core.models import FutureBudget, FutureBudgetKnownChange, FutureBudgetPersonnel

ZERO = Decimal("0.00")
SHEET_NAMES = [
    "01 Cover & Parameters", "02 Executive Summary", "03 Existing Budget vs Actual",
    "04 Head-wise Analysis", "05 System Suggested", "06 Management Proposed",
    "07 Approved Budget", "08 Management Requirements", "09 Unit-wise Analysis",
    "10 Internal Section Analysis", "11 Unit-Section Analysis", "12 Personnel Effect",
    "13 Known Changes", "14 One-Time - Discontinued", "15 Variance & Decision Flags",
    "16 Audit & Approval",
]


def _money(value):
    if value is None:
        return None
    return (value or ZERO).quantize(Decimal("0.01"))


def _diff(left, right):
    if left is None or right is None:
        return None
    return _money(left - right)


def _excel_value(value):
    """Convert timezone-aware datetimes to Excel-safe naive values."""
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo is not None else value
    if isinstance(value, date):
        return value
    return value


def _sheet(ws, title, headers, rows, filters):
    ws.title = title
    ws.append(["ABMS — Enterprise Accounts & Business Management System (EABMS)"])
    ws.append([title])
    ws.append([f"Account: {filters['account']}"])
    ws.append([f"Source Financial Year: {filters['source_fy']}"])
    ws.append([f"Target Financial Year: {filters['target_fy']}"])
    ws.append([f"Unit Filter: {', '.join(filters['units']) if filters['units'] else 'All'}"])
    ws.append([f"Internal Section Filter: {', '.join(filters['sections']) if filters['sections'] else 'All'}"])
    ws.append([])
    ws.append(headers)
    for row in rows:
        ws.append([_excel_value(value) for value in row])
    header_row = 9
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A10"
    if ws.max_row > header_row:
        ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"
        safe = "Tbl" + "".join(ch for ch in title.title() if ch.isalnum())[:20]
        table = Table(displayName=safe, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    for column in ws.columns:
        width = min(max(len(str(c.value or "")) for c in column) + 2, 55)
        ws.column_dimensions[get_column_letter(column[0].column)].width = max(width, 10)


def export_future_budget_management_workbook(account, source_fy, target_fy, units=None, internal_sections=None):
    """Build the approved 16-sheet management workbook from ABMS SSOT data."""
    units = list(units or [])
    internal_sections = list(internal_sections or [])
    base = export_budget_planning_workbook(account, source_fy, target_fy, units=units or None, internal_sections=internal_sections or None)
    future = FutureBudget.objects.filter(account=account, source_fy=source_fy, target_fy=target_fy).first()
    if not future:
        return base
    wb = load_workbook(BytesIO(base))
    for name in list(wb.sheetnames):
        del wb[name]
    filters = {"account": getattr(account, "name", "All"), "source_fy": getattr(source_fy, "name", "All"), "target_fy": getattr(target_fy, "name", "All"), "units": units, "sections": internal_sections}
    summary = get_future_budget_summary(future)
    lines = list(future.lines.select_related("account_head").order_by("account_head__name"))
    ws = wb.create_sheet(SHEET_NAMES[0], 0)
    _sheet(ws, SHEET_NAMES[0], ["Parameter", "Value"], [["Account", filters["account"]], ["Source FY", filters["source_fy"]], ["Target FY", filters["target_fy"]], ["Units", ", ".join(units) if units else "All"], ["Internal Sections", ", ".join(internal_sections) if internal_sections else "All"], ["Future Budget ID", future.id], ["Future Budget Status", future.get_status_display() if hasattr(future, "get_status_display") else future.status], ["Line Count", summary["line_count"]]], filters)
    ws = wb.create_sheet(SHEET_NAMES[1], 1)
    _sheet(ws, SHEET_NAMES[1], ["Metric", "Amount"], [["Previous Budget", _money(summary["previous_budget_total"])], ["Previous Actual", _money(summary["previous_actual_total"])], ["Historical Additional Actual", _money(summary["additional_total"])], ["Routine", _money(summary["routine_total"])], ["Management Continuing", _money(summary["management_continuing_total"])], ["Management Additional Requirement", _money(summary["management_additional_total"])], ["Known Changes", _money(summary["known_changes_total"])], ["Personnel Effect", _money(summary["personnel_effect_total"])], ["One-Time", _money(summary["one_time_total"])], ["System Suggested", _money(summary["system_suggested_total"])], ["Management Proposed", _money(summary["user_proposed_total"])], ["Approved Budget", _money(summary["approved_total"])], ["Proposed vs Previous Actual", _money(summary["diff_from_prev_actual"])], ["Proposed vs Previous Budget", _money(summary["diff_from_prev_budget"])]], filters)
    head_rows = [[l.account_head.name, _money(l.previous_budget), _money(l.previous_actual), _money(l.addl_amount), _money(l.continuing_amount), _money(l.known_change_amount), _money(l.personnel_amount), _money(l.onetime_amount), _money(l.system_suggested), _money(l.user_proposed), _money(l.approved_amount), l.continuing_flag or ""] for l in lines]
    _sheet(wb.create_sheet(SHEET_NAMES[2], 2), SHEET_NAMES[2], ["Account Head", "Previous Budget", "Previous Actual", "Variance"], [[r[0], r[1], r[2], _diff(r[2], r[1])] for r in head_rows], filters)
    _sheet(wb.create_sheet(SHEET_NAMES[3], 3), SHEET_NAMES[3], ["Account Head", "Previous Budget", "Previous Actual", "Historical Additional", "Continuing", "Known Changes", "Personnel", "One-Time", "System Suggested", "Management Proposed", "Approved", "Flag"], head_rows, filters)
    _sheet(wb.create_sheet(SHEET_NAMES[4], 4), SHEET_NAMES[4], ["Account Head", "Previous Actual", "Routine", "Continuing", "Known Changes", "Personnel", "One-Time", "Management Additional", "System Suggested"], [[l.account_head.name, _money(l.previous_actual), _money(l.routine_amount), _money(l.continuing_amount), _money(l.known_change_amount), _money(l.personnel_amount), _money(l.onetime_amount), _money(l.suggestion_breakdown.get("management_additional_requirement", 0)), _money(l.system_suggested)] for l in lines], filters)
    _sheet(wb.create_sheet(SHEET_NAMES[5], 5), SHEET_NAMES[5], ["Account Head", "System Suggested", "Management Proposed", "Difference", "Decision"], [[l.account_head.name, _money(l.system_suggested), _money(l.user_proposed), _diff(l.user_proposed, l.system_suggested), "Not Entered" if l.user_proposed is None else ("Unchanged" if l.user_proposed == l.system_suggested else "Management Override")] for l in lines], filters)
    _sheet(wb.create_sheet(SHEET_NAMES[6], 6), SHEET_NAMES[6], ["Account Head", "Management Proposed", "Approved", "Difference", "Status"], [[l.account_head.name, _money(l.user_proposed), _money(l.approved_amount), _diff(l.approved_amount, l.user_proposed), "Approved" if l.approved_amount is not None else "Pending"] for l in lines], filters)
    inputs = FutureBudgetManagementInput.objects.filter(future_budget=future).select_related("account_head").prefetch_related("units", "internal_sections").order_by("id")
    input_rows = []
    for r in inputs:
        us = list(r.units.values_list("name", flat=True)); ss = list(r.internal_sections.values_list("name", flat=True))
        if units and us and not set(us).intersection(units): continue
        if internal_sections and ss and not set(ss).intersection(internal_sections): continue
        input_rows.append([r.id, r.account_head.name, r.input_type, r.status, _money(r.amount), ", ".join(us) or "All", ", ".join(ss) or "All", r.description, r.reference])
    _sheet(wb.create_sheet(SHEET_NAMES[7], 7), SHEET_NAMES[7], ["ID", "Account Head", "Type", "Status", "Amount", "Units", "Internal Sections", "Description", "Reference"], input_rows, filters)
    org = build_future_organisational_analysis(future, units=units or None, internal_sections=internal_sections or None)
    _sheet(wb.create_sheet(SHEET_NAMES[8], 8), SHEET_NAMES[8], ["Unit", "Previous Actual", "Future Allocated", "Difference", "Change %"], [[r["unit"], _money(r["previous_actual"]), _money(r["future_allocated"]), _money(r["difference"]), r["change_pct"]] for r in org["unit_rows"]], filters)
    _sheet(wb.create_sheet(SHEET_NAMES[9], 9), SHEET_NAMES[9], ["Internal Section", "Previous Actual", "Future Allocated", "Difference", "Change %"], [[r["internal_section"], _money(r["previous_actual"]), _money(r["future_allocated"]), _money(r["difference"]), r["change_pct"]] for r in org["section_rows"]], filters)
    _sheet(wb.create_sheet(SHEET_NAMES[10], 10), SHEET_NAMES[10], ["Unit", "Internal Section", "Future Allocated"], [[r["unit"], r["internal_section"], _money(r["future_allocated"])] for r in org["matrix_rows"]], filters)
    personnel = FutureBudgetPersonnel.objects.filter(future_budget=future).order_by("id")
    _sheet(wb.create_sheet(SHEET_NAMES[11], 11), SHEET_NAMES[11], ["ID", "Category", "Current Annual Cost", "Projected Cost", "Effect"], [[r.id, r.category_label, _money(r.current_annual_cost), _money(r.projected_cost), _money(r.projected_cost - r.current_annual_cost)] for r in personnel], filters)
    changes = FutureBudgetKnownChange.objects.filter(future_budget=future).select_related("account_head").order_by("id")
    _sheet(wb.create_sheet(SHEET_NAMES[12], 12), SHEET_NAMES[12], ["ID", "Account Head", "Change Type", "Amount", "Description", "Reference"], [[r.id, r.account_head.name if r.account_head else "Unassigned", r.change_type, _money(r.amount), r.description, r.reference] for r in changes], filters)
    _sheet(wb.create_sheet(SHEET_NAMES[13], 13), SHEET_NAMES[13], ["Account Head", "One-Time Amount", "System Deduction"], [[l.account_head.name, _money(l.onetime_amount), _money(l.onetime_amount)] for l in lines if l.onetime_amount], filters)
    proposal_overrides = sum(1 for l in lines if l.user_proposed is not None and l.user_proposed != l.system_suggested)
    proposals_missing = sum(1 for l in lines if l.user_proposed is None)
    approvals_missing = sum(1 for l in lines if l.approved_amount is None)
    _sheet(wb.create_sheet(SHEET_NAMES[14], 14), SHEET_NAMES[14], ["Account Head", "Previous Budget", "Previous Actual", "System Suggested", "Management Proposed", "Approved", "Proposal vs System", "Proposal vs Previous", "Decision Flag"], [[l.account_head.name, _money(l.previous_budget), _money(l.previous_actual), _money(l.system_suggested), _money(l.user_proposed), _money(l.approved_amount), _diff(l.user_proposed, l.system_suggested), _diff(l.user_proposed, l.previous_actual), "NOT ENTERED" if l.user_proposed is None else ("REVIEW" if l.user_proposed != l.system_suggested else "OK")] for l in lines], filters)
    _sheet(wb.create_sheet(SHEET_NAMES[15], 15), SHEET_NAMES[15], ["Control", "Value"], [["Future Budget ID", future.id], ["Status", future.get_status_display() if hasattr(future, "get_status_display") else future.status], ["Line Count", summary["line_count"]], ["Lines Without Management Proposed", proposals_missing], ["Lines Without Approved Amount", approvals_missing], ["Management Proposal Overrides", proposal_overrides], ["System Suggested Total", _money(summary["system_suggested_total"])], ["Management Proposed Total", _money(summary["user_proposed_total"])], ["Approved Total", _money(summary["approved_total"])], ["Created By", getattr(getattr(future, "created_by", None), "username", "")], ["Created At", getattr(future, "created_at", None)], ["Updated By", getattr(getattr(future, "updated_by", None), "username", "")], ["Updated At", getattr(future, "updated_at", None)], ["Reference", future.reference], ["Remarks", future.remarks], ["SSOT", "ABMS database; Excel is an export/management working paper"]], filters)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
