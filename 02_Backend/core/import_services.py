import openpyxl
import csv
from decimal import Decimal, InvalidOperation
from datetime import datetime
from django.db import transaction
from django.core.exceptions import ValidationError

from core.models import (
    ImportBatch, ImportedRecord, Transaction, AccountHead,
    FinancialYear, Account
)


def _read_tabular_rows(file_path):
    """Read XLSX, legacy XLS (Excel 2003), or CSV into normalized row dictionaries."""
    ext = file_path.lower().rsplit('.', 1)[-1]
    rows = []

    if ext == 'xlsx':
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue
            rows.append({'original_row_number': row_idx, 'row_data': dict(zip(headers, row))})
        wb.close()
        return rows

    if ext == 'xls':
        import xlrd
        workbook = xlrd.open_workbook(file_path, on_demand=True)
        sheet = workbook.sheet_by_index(0)
        headers = [sheet.cell_value(0, col) for col in range(sheet.ncols)]
        for row_idx in range(1, sheet.nrows):
            values = [sheet.cell_value(row_idx, col) for col in range(sheet.ncols)]
            if all(value in ('', None) for value in values):
                continue
            rows.append({'original_row_number': row_idx + 1, 'row_data': dict(zip(headers, values))})
        workbook.release_resources()
        return rows

    if ext == 'csv':
        with open(file_path, newline='', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            for row_idx, row in enumerate(reader, start=2):
                if not any(str(value or '').strip() for value in row.values()):
                    continue
                rows.append({'original_row_number': row_idx, 'row_data': dict(row)})
        return rows

    raise ValueError("Unsupported file format. Supported formats are .xlsx, .xls (Excel 2003), and .csv")


def parse_import_file(batch):
    """Parse uploaded XLSX/XLS/CSV and create ImportedRecord objects."""
    rows = _read_tabular_rows(batch.source_file.path)
    records = [ImportedRecord(batch=batch, original_row_number=r['original_row_number'], row_data=r['row_data'], validation_status=ImportedRecord.Status.INVALID, error_message='Not validated yet') for r in rows]
    ImportedRecord.objects.bulk_create(records)
    batch.total_rows = len(records)
    batch.status = ImportBatch.Status.DRAFT
    batch.save()


def apply_ngl_head_mapping(raw_head):
    if not raw_head:
        return None
    raw = str(raw_head).strip()
    manual_map = {
        'Salaries': 'Salaries',
        'Contract Salaries': 'Salaries',
        'Misc Exp': 'Office Misc Exp',
        'Fin Asst Indis': 'Fin Asst Indls',
        'Amount Tfr from HQO': 'Received from HQO',
        'Capacity Enhancement Project': 'Capacity Enhancement Project',
        'Pass Through Transactions': 'All pass through transactions',
        'cafeteria': 'cafeteria',
        'Amount Tfr from EVT': 'Received from Ceremony Fund Acct',
        'Amount Tfr from TRV': 'Received from Umra & Haj Acct',
    }
    return manual_map.get(raw, raw)


def apply_unit_mapping(raw_unit):
    if not raw_unit:
        return {}
    raw = str(raw_unit).strip().lower()
    if raw == 'addl':
        return {'outflow_nature': 'Addl'}
    elif raw == 'routine':
        return {'outflow_nature': 'Routine'}
    elif raw == 'reimbursement':
        return {'reimbursement_flag': True}
    elif raw == 'pass through':
        return {'pass_through': True}
    return {}


def validate_import_batch(batch, mapping):
    valid_count = failed_count = duplicate_count = 0
    records = ImportedRecord.objects.filter(batch=batch)
    for record in records:
        data = record.row_data
        errors = []
        try:
            date_val = data.get(mapping.get('date'))
            amount_val = data.get(mapping.get('amount'))
            head_code_val = data.get(mapping.get('head_code'))
            ref_val = data.get(mapping.get('reference'))
            desc_val = data.get(mapping.get('description'))
            dir_val = data.get(mapping.get('direction', ''))
            if not date_val or not amount_val or not head_code_val:
                errors.append('Missing mandatory fields (Date, Amount, HeadCode).')
            parsed_date = None
            if date_val:
                if isinstance(date_val, datetime):
                    parsed_date = date_val.date()
                else:
                    try:
                        parsed_date = datetime.strptime(str(date_val).strip(), '%Y-%m-%d').date()
                    except ValueError:
                        errors.append(f'Invalid date format: {date_val}')
            parsed_amount = None
            if amount_val:
                try:
                    parsed_amount = Decimal(str(amount_val).replace(',', ''))
                except InvalidOperation:
                    errors.append(f'Invalid amount format: {amount_val}')
            head = None
            if head_code_val:
                mapped_name = apply_ngl_head_mapping(head_code_val)
                head = AccountHead.objects.filter(account=batch.account, name__iexact=mapped_name).first()
                if not head:
                    errors.append(f'Unknown Account Head: {head_code_val} (mapped to {mapped_name})')
            fy = None
            if parsed_date:
                fy = FinancialYear.objects.filter(start_date__lte=parsed_date, end_date__gte=parsed_date).first()
                if not fy:
                    errors.append(f'No Financial Year covers date {parsed_date}')
            direction = Transaction.Direction.OUTFLOW
            if dir_val and str(dir_val).upper().strip() in ['INFLOW', 'RECEIPT']:
                direction = Transaction.Direction.INFLOW
            if parsed_date and parsed_amount and head:
                is_duplicate = Transaction.objects.filter(account=batch.account, transaction_date=parsed_date, amount=parsed_amount, account_head=head).exists()
                if is_duplicate:
                    errors.append('Potential duplicate found.')
                    record.duplicate_status = True
        except Exception as e:
            errors.append(f'Error parsing row: {str(e)}')
        if errors:
            record.validation_status = ImportedRecord.Status.INVALID
            if record.duplicate_status:
                record.validation_status = ImportedRecord.Status.DUPLICATE
                duplicate_count += 1
            else:
                failed_count += 1
            record.error_message = '; '.join(errors)
        else:
            record.validation_status = ImportedRecord.Status.VALID
            record.error_message = ''
            valid_count += 1
        record.save()
    batch.status = ImportBatch.Status.PREVIEWED
    batch.failed_rows = failed_count
    batch.duplicate_rows = duplicate_count
    batch.save()


def execute_import_batch(batch, mapping):
    """Commits valid ImportedRecords to the Transaction table safely."""
    if batch.status != ImportBatch.Status.PREVIEWED:
        raise ValidationError('Batch must be previewed before importing.')
    valid_records = batch.records.filter(validation_status=ImportedRecord.Status.VALID)
    imported_count = 0
    with transaction.atomic():
        for record in valid_records:
            data = record.row_data
            date_val = data.get(mapping.get('date'))
            parsed_date = date_val.date() if isinstance(date_val, datetime) else datetime.strptime(str(date_val).strip(), '%Y-%m-%d').date()
            amount_val = Decimal(str(data.get(mapping.get('amount'))).replace(',', ''))
            head_code_val = str(data.get(mapping.get('head_code'))).strip()
            head = AccountHead.objects.get(account=batch.account, name__iexact=apply_ngl_head_mapping(head_code_val))
            fy = FinancialYear.objects.get(start_date__lte=parsed_date, end_date__gte=parsed_date)
            unit_val = data.get(mapping.get('unit_type', ''))
            unit_map = apply_unit_mapping(unit_val)
            dir_val = data.get(mapping.get('direction', ''))
            direction = Transaction.Direction.INFLOW if str(dir_val).upper().strip() in ['INFLOW', 'RECEIPT'] else Transaction.Direction.OUTFLOW
            if unit_map.get('reimbursement_flag'):
                direction = Transaction.Direction.INFLOW
            pass_through = unit_map.get('pass_through', False)
            outflow_nature_obj = None
            if unit_map.get('outflow_nature'):
                from core.models import OutflowNature
                outflow_nature_obj = OutflowNature.objects.filter(code__iexact=unit_map['outflow_nature']).first()
            desc_val = data.get(mapping.get('description')) or ''
            ref_val = data.get(mapping.get('reference')) or ''
            txn = Transaction(account=batch.account, financial_year=fy, account_head=head, transaction_date=parsed_date, direction=direction, amount=amount_val, payee_or_source='Imported Record', remarks=desc_val, reference=ref_val, pass_through=pass_through, outflow_nature=outflow_nature_obj, status=Transaction.Status.ACTIVE, created_by=batch.imported_by, is_historical=True)
            txn.save()
            record.target_transaction = txn
            record.validation_status = ImportedRecord.Status.IMPORTED
            record.save()
            imported_count += 1
        batch.status = ImportBatch.Status.COMPLETED
        batch.imported_rows = imported_count
        batch.save()
