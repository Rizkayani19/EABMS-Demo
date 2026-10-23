import io
from decimal import Decimal
import docx
import pytest
from django.test import Client
from django.urls import reverse
from django.contrib.auth import get_user_model

from core.models import (
    FutureBudget, FutureBudgetLine, Account, FinancialYear,
    Installation, AccountHead, Transaction
)
from core.budget_word_report import generate_budget_word_report


@pytest.fixture
def sample_budget(db):
    user = get_user_model().objects.create_superuser('admin_user', 'admin@example.com', 'pass123')
    inst = Installation.objects.create(installation_code='INST1', installation_name='Installation 1')
    acct = Account.objects.create(installation=inst, code='NGL', name='NexGen Holdings')
    fy_src = FinancialYear.objects.create(code='FY-25-26', name='FY 2025-26', start_date='2025-07-01', end_date='2026-06-30')
    fy_tgt = FinancialYear.objects.create(code='FY-26-27', name='FY 2026-27', start_date='2026-07-01', end_date='2027-06-30')

    fb = FutureBudget.objects.create(
        account=acct,
        source_fy=fy_src,
        target_fy=fy_tgt,
        status=FutureBudget.Status.APPROVED,
        created_by=user,
    )

    head_cap = AccountHead.objects.create(account=acct, code='CAP', name='Capacity Enhancement Project', head_type='OUTFLOW')
    head_sal = AccountHead.objects.create(account=acct, code='SAL', name='Salaries', head_type='OUTFLOW')

    FutureBudgetLine.objects.create(
        future_budget=fb,
        account_head=head_cap,
        previous_budget=Decimal('50000000.00'),
        previous_actual=Decimal('70000000.00'),
        routine_amount=Decimal('30000000.00'),
        system_suggested=Decimal('35000000.00'),
        approved_amount=Decimal('32000000.00'),
    )
    FutureBudgetLine.objects.create(
        future_budget=fb,
        account_head=head_sal,
        previous_budget=Decimal('28000000.00'),
        previous_actual=Decimal('48000000.00'),
        routine_amount=Decimal('32000000.00'),
        system_suggested=Decimal('46000000.00'),
        approved_amount=Decimal('46000000.00'),
    )

    Transaction.objects.create(
        account=acct,
        financial_year=fy_src,
        account_head=head_cap,
        amount=Decimal('5682300.00'),
        direction=Transaction.Direction.OUTFLOW,
        status=Transaction.Status.ACTIVE,
        pass_through=True,
        transaction_date='2025-07-02',
        description='Paid to Finance on acct of training course',
        remarks='Re-imbursed by Group',
    )

    return fb, user


def test_budget_word_report_generation(sample_budget):
    fb, _ = sample_budget
    data = generate_budget_word_report(fb)
    assert isinstance(data, bytes)
    assert len(data) > 1000

    doc = docx.Document(io.BytesIO(data))
    assert doc.sections[0].header.paragraphs[0].text == 'CONFIDENTIAL'
    assert doc.sections[0].footer.paragraphs[0].text == 'CONFIDENTIAL'

    all_text = " ".join(p.text for p in doc.paragraphs)
    assert 'Approved by' in all_text
    assert 'NEXGEN HOLDINGS' in all_text
    assert 'Overview' in all_text
    assert 'Capacity Enhancement Project' in all_text
    assert 'Salaries' in all_text
    assert 'Recoms' in all_text
    assert len(doc.tables) >= 6


def test_future_budget_word_export_view(sample_budget):
    fb, user = sample_budget

    client = Client()
    client.force_login(user)

    session = client.session
    session['active_account_id'] = fb.account.pk
    session['active_fy_id'] = fb.source_fy.pk
    session.save()

    url = reverse('core:future_budget_word', kwargs={'pk': fb.pk})
    resp = client.get(url)
    assert resp.status_code == 200
    assert resp['Content-Type'] == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert 'attachment; filename=' in resp['Content-Disposition']
    assert len(resp.content) > 1000

