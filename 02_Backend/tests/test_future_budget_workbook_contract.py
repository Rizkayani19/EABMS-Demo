"""Acceptance tests for the generated Future Budget management workbook."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from core.future_budget_management_export import SHEET_NAMES, export_future_budget_management_workbook


class FakeQuerySet(list):
    def select_related(self, *args, **kwargs): return self
    def prefetch_related(self, *args, **kwargs): return self
    def order_by(self, *args, **kwargs): return self
    def first(self): return self[0] if self else None


class FakeManyToMany:
    def __init__(self, names): self.names = names
    def values_list(self, *args, **kwargs): return self.names


class FutureBudgetWorkbookContractTests(SimpleTestCase):
    def _base_workbook_bytes(self):
        wb = Workbook()
        wb.active.title = "Base Workbook"
        output = BytesIO()
        wb.save(output)
        return output.getvalue()

    def _fixture(self):
        account = SimpleNamespace(name="NexGen Holdings")
        source_fy = SimpleNamespace(name="2025-26")
        target_fy = SimpleNamespace(name="2026-27")
        head = SimpleNamespace(name="Contract Salaries")
        line = SimpleNamespace(
            account_head=head, previous_budget=Decimal("1000000.00"), previous_actual=Decimal("900000.00"),
            addl_amount=Decimal("50000.00"), continuing_amount=Decimal("25000.00"), known_change_amount=Decimal("100000.00"),
            personnel_amount=Decimal("75000.00"), onetime_amount=Decimal("20000.00"), system_suggested=Decimal("1080000.00"),
            user_proposed=Decimal("1150000.00"), approved_amount=Decimal("1125000.00"), continuing_flag="needs_management_review",
            routine_amount=Decimal("900000.00"), suggestion_breakdown={"management_additional_requirement": Decimal("50000.00")},
        )
        future = SimpleNamespace(
            id=42, status="DRAFT", reference="FB-26-27-001", remarks="Management planning working paper",
            created_by=SimpleNamespace(username="admin"), updated_by=SimpleNamespace(username="management"),
            created_at=datetime(2026, 8, 23, 18, 0, 0), updated_at=datetime(2026, 8, 23, 18, 30, 0),
            lines=FakeQuerySet([line]), get_status_display=lambda: "Draft",
        )
        management_input = SimpleNamespace(
            id=7, account_head=head, input_type="ADDITIONAL", status="ACTIVE", amount=Decimal("50000.00"),
            description="Additional requirement", reference="MR-001", units=FakeManyToMany(["Unit A", "Unit B"]),
            internal_sections=FakeManyToMany(["Section 1"]),
        )
        personnel = SimpleNamespace(id=8, category_label="Contract Salaries", current_annual_cost=Decimal("900000.00"), projected_cost=Decimal("975000.00"))
        known_change = SimpleNamespace(id=9, account_head=head, change_type="NEW", amount=Decimal("100000.00"), description="Known new requirement", reference="KC-001")
        summary = {
            "previous_budget_total": Decimal("1000000.00"), "previous_actual_total": Decimal("900000.00"), "additional_total": Decimal("50000.00"),
            "routine_total": Decimal("900000.00"), "management_continuing_total": Decimal("25000.00"), "management_additional_total": Decimal("50000.00"),
            "management_other_total": Decimal("0.00"), "known_changes_total": Decimal("100000.00"), "personnel_effect_total": Decimal("75000.00"),
            "one_time_total": Decimal("20000.00"), "system_suggested_total": Decimal("1080000.00"), "user_proposed_total": Decimal("1150000.00"),
            "approved_total": Decimal("1125000.00"), "diff_from_prev_actual": Decimal("250000.00"), "diff_from_prev_budget": Decimal("150000.00"), "line_count": 1,
        }
        organisation = {
            "unit_rows": [{"unit": "Unit A", "previous_actual": Decimal("100.00"), "future_allocated": Decimal("120.00"), "difference": Decimal("20.00"), "change_pct": "20.00%"}],
            "section_rows": [{"internal_section": "Section 1", "previous_actual": Decimal("100.00"), "future_allocated": Decimal("120.00"), "difference": Decimal("20.00"), "change_pct": "20.00%"}],
            "matrix_rows": [{"unit": "Unit A", "internal_section": "Section 1", "future_allocated": Decimal("120.00")}],
        }
        return account, source_fy, target_fy, future, line, management_input, personnel, known_change, summary, organisation

    def test_required_management_sheets_are_defined_in_order(self):
        self.assertEqual(SHEET_NAMES, [
            "01 Cover & Parameters", "02 Executive Summary", "03 Existing Budget vs Actual", "04 Head-wise Analysis",
            "05 System Suggested", "06 Management Proposed", "07 Approved Budget", "08 Management Requirements",
            "09 Unit-wise Analysis", "10 Internal Section Analysis", "11 Unit-Section Analysis", "12 Personnel Effect",
            "13 Known Changes", "14 One-Time - Discontinued", "15 Variance & Decision Flags", "16 Audit & Approval",
        ])

    @patch("core.future_budget_management_export.build_future_organisational_analysis")
    @patch("core.future_budget_management_export.FutureBudgetKnownChange")
    @patch("core.future_budget_management_export.FutureBudgetPersonnel")
    @patch("core.future_budget_management_export.FutureBudgetManagementInput")
    @patch("core.future_budget_management_export.get_future_budget_summary")
    @patch("core.future_budget_management_export.FutureBudget")
    @patch("core.future_budget_management_export.export_budget_planning_workbook")
    def test_generated_workbook_has_contract_filters_and_decision_values(self, base_export, future_budget_model, summary_service, management_input_model, personnel_model, known_change_model, organisation_service):
        account, source_fy, target_fy, future, line, management_input, personnel, known_change, summary, organisation = self._fixture()
        base_export.return_value = self._base_workbook_bytes()
        future_budget_model.objects.filter.return_value = FakeQuerySet([future])
        summary_service.return_value = summary
        management_input_model.objects.filter.return_value = FakeQuerySet([management_input])
        personnel_model.objects.filter.return_value = FakeQuerySet([personnel])
        known_change_model.objects.filter.return_value = FakeQuerySet([known_change])
        organisation_service.return_value = organisation
        wb = load_workbook(BytesIO(export_future_budget_management_workbook(account, source_fy, target_fy, units=["Unit A", "Unit B"], internal_sections=["Section 1"])))
        self.assertEqual(wb.sheetnames, SHEET_NAMES)
        proposed = wb["06 Management Proposed"][10]
        self.assertEqual(proposed[0].value, "Contract Salaries"); self.assertEqual(proposed[1].value, 1080000); self.assertEqual(proposed[2].value, 1150000); self.assertEqual(proposed[3].value, 70000); self.assertEqual(proposed[4].value, "Management Override")
        approved = wb["07 Approved Budget"][10]
        self.assertEqual(approved[0].value, "Contract Salaries"); self.assertEqual(approved[1].value, 1150000); self.assertEqual(approved[2].value, 1125000); self.assertEqual(approved[3].value, -25000); self.assertEqual(approved[4].value, "Approved")
        requirements = wb["08 Management Requirements"][10]
        self.assertEqual(requirements[5].value, "Unit A, Unit B"); self.assertEqual(requirements[6].value, "Section 1")
        personnel_row = wb["12 Personnel Effect"][10]
        self.assertEqual(personnel_row[0].value, 8); self.assertEqual(personnel_row[1].value, "Contract Salaries"); self.assertEqual(personnel_row[4].value, 75000)

    @patch("core.future_budget_management_export.build_future_organisational_analysis", return_value={"unit_rows": [], "section_rows": [], "matrix_rows": []})
    @patch("core.future_budget_management_export.FutureBudgetKnownChange")
    @patch("core.future_budget_management_export.FutureBudgetPersonnel")
    @patch("core.future_budget_management_export.FutureBudgetManagementInput")
    @patch("core.future_budget_management_export.get_future_budget_summary")
    @patch("core.future_budget_management_export.FutureBudget")
    @patch("core.future_budget_management_export.export_budget_planning_workbook")
    def test_missing_proposal_is_preserved_as_blank_and_does_not_crash(self, base_export, future_budget_model, summary_service, management_input_model, personnel_model, known_change_model, organisation_service):
        account, source_fy, target_fy, future, line, management_input, personnel, known_change, summary, organisation = self._fixture()
        line.user_proposed = None
        line.approved_amount = None
        summary["user_proposed_total"] = Decimal("0.00")
        summary["approved_total"] = Decimal("0.00")
        base_export.return_value = self._base_workbook_bytes()
        future_budget_model.objects.filter.return_value = FakeQuerySet([future])
        summary_service.return_value = summary
        management_input_model.objects.filter.return_value = FakeQuerySet([management_input])
        personnel_model.objects.filter.return_value = FakeQuerySet([personnel])
        known_change_model.objects.filter.return_value = FakeQuerySet([known_change])
        organisation_service.return_value = organisation
        wb = load_workbook(BytesIO(export_future_budget_management_workbook(account, source_fy, target_fy)))
        proposed = wb["06 Management Proposed"][10]
        self.assertIsNone(proposed[2].value); self.assertIsNone(proposed[3].value); self.assertEqual(proposed[4].value, "Not Entered")
        approved = wb["07 Approved Budget"][10]
        self.assertIsNone(approved[1].value); self.assertIsNone(approved[2].value); self.assertIsNone(approved[3].value); self.assertEqual(approved[4].value, "Pending")
        flags = wb["15 Variance & Decision Flags"][10]
        self.assertEqual(flags[8].value, "NOT ENTERED")
